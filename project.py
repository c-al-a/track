import streamlit as st
import pandas as pd
from openpyxl import load_workbook

FILE_PATH = "tracker.xlsx"

# =========================
# LOAD DATA
# =========================
@st.cache_data
def load_data():
    movement = pd.read_excel(FILE_PATH, sheet_name="MOVEMENT LOG")
    lifecycle = pd.read_excel(FILE_PATH, sheet_name="NEW HIRE_PROMOTION_RESIGNATION")
    resources = pd.read_excel(FILE_PATH, sheet_name="RESOURCES")
    leaves = pd.read_excel(FILE_PATH, sheet_name="LEAVES")

    return movement, lifecycle, resources, leaves


def load_named_range(name):
    wb = load_workbook(FILE_PATH, data_only=True)
    rng = wb.defined_names[name]

    values = []
    for sheet, coord in rng.destinations:
        ws = wb[sheet]
        for row in ws[coord]:
            if isinstance(row, tuple):
                values.extend([c.value for c in row])
            else:
                values.append(row.value)

    return values


@st.cache_data
def load_lists():
    names = load_named_range("NameList")
    holidays = set(pd.to_datetime(load_named_range("Holidays")).date)
    return names, holidays


# =========================
# CLEAN VALUES
# =========================
def clean_team(val):
    if val is None or pd.isna(val) or str(val).strip() == "":
        return "Blank"
    return str(val)


# =========================
# PREPARE MOVEMENT
# =========================
def prepare_movement(df):
    df = df.copy()

    df["START DATE"] = pd.to_datetime(df["START DATE"])
    df["END DATE"] = pd.to_datetime(df["END DATE"])

    # open-ended TEMP logs
    df["END DATE"] = df["END DATE"].fillna(pd.Timestamp("2100-01-01"))

    df["DESTINATION TEAM/CLIENT"] = df["DESTINATION TEAM/CLIENT"].apply(clean_team)

    return df

# =========================
# PREPARE LEAVES
# =========================
def prepare_leaves(df):
    df = df.copy()

    df["DATE"] = pd.to_datetime(df["DATE"])

    return df


# =========================
# SPLIT CORE / TEMP
# =========================
def split_core_temp(df):
    core = df[df["ASSIGNMENT TYPE"] == "PERMANENT"].copy()
    temp = df[df["ASSIGNMENT TYPE"] == "TEMPORARY"].copy()
    return core, temp
    
def get_all_core_teams(movement):
    movement = prepare_movement(movement)
    core_df, _ = split_core_temp(movement)

    return sorted(core_df["DESTINATION TEAM/CLIENT"].dropna().unique())


# =========================
# BUILD TEMP INTERVALS
# =========================
def build_temp_intervals(temp_rows):
    temp_rows = temp_rows.sort_values("START DATE")

    intervals = []

    for _, row in temp_rows.iterrows():
        start = row["START DATE"]
        end = row["END DATE"]
        team = clean_team(row["DESTINATION TEAM/CLIENT"])

        next_rows = temp_rows[temp_rows["START DATE"] > start]

        if not next_rows.empty:
            next_start = next_rows.iloc[0]["START DATE"]
            end = min(end, next_start)

        intervals.append((start, end, team))

    return intervals

# =========================
# CHECK LEAVE
# =========================
def is_on_leave(name, date, leaves):
    leaves = prepare_leaves(leaves)

    emp_leaves = leaves[leaves["NAME"] == name]

    leave_dates = set(emp_leaves["DATE"].dt.date)

    return date.date() in leave_dates


# =========================
# RESOLVE EMPLOYEE TIMELINE
# =========================
def resolve_employee(name, movement, timeline, holidays, leaves):
    movement = prepare_movement(movement)
    core_df, temp_df = split_core_temp(movement)

    core_rows = core_df[core_df["NAME"] == name]
    temp_rows = temp_df[temp_df["NAME"] == name]

    if core_rows.empty:
        return ["Blank"] * len(timeline)

    core_team = clean_team(core_rows.iloc[-1]["DESTINATION TEAM/CLIENT"])

    temp_intervals = build_temp_intervals(temp_rows)

    result = []

    for t in timeline:
        active = None

        for start, end, team in temp_intervals:
            if start <= t <= end:
                active = team
                break

                # LEAVE
        if is_on_leave(name, t, leaves):
            result.append("LEAVE")

        # WEEKENDS / HOLIDAYS
        elif t.weekday() >= 5 or t.date() in holidays:
            result.append("NON-WORK DAY")

        # TEMP ASSIGNMENT
        elif active:
            result.append(active)

        # CORE TEAM
        else:
            result.append(core_team)

    return result


# =========================
# GET CORE TEAM ONLY
# =========================
def get_core_team(name, movement):
    movement = prepare_movement(movement)
    core_df, _ = split_core_temp(movement)

    core_rows = core_df[core_df["NAME"] == name]

    if core_rows.empty:
        return None

    return clean_team(core_rows.iloc[-1]["DESTINATION TEAM/CLIENT"])


# =========================
# TIMELINE
# =========================
def build_timeline(start_date, end_date):
    return pd.date_range(start_date, end_date, freq="D")


# =========================
# COLORS
# =========================
def generate_color_map(values):

    color_map = {
        "NON-WORK DAY": "#555555",
        "LEAVE": "#444444",
        "Blank": "#ffffff"
    }

    palette = [
        "#dbeafe",
        "#ffedd5",
        "#dcfce7",
        "#fee2e2",
        "#ede9fe",
        "#f3e8e2",
        "#fce7f3",
        "#e5e7eb",
        "#fef9c3",
        "#cffafe",
        "#fde68a",
        "#ddd6fe",
        "#bfdbfe",
        "#fecaca",
        "#bbf7d0"
    ]

    unique = sorted(set(values))

    team_values = [
        v for v in unique
        if v not in color_map
    ]

    for i, v in enumerate(team_values):
        color_map[v] = palette[i % len(palette)]

    return color_map


# =========================
# HTML TABLE
# =========================
def render_colored_table(df, color_map):
    html = """
    <style>
        table { border-collapse: collapse; font-size: 12px; }
        td, th { border: 1px solid #ddd; padding: 4px; text-align: center; }
    </style>
    <table>
    """

    html += "<tr><th>Employee</th>"

    for col in df.columns:
        html += f"<th>{col.strftime('%Y-%m-%d')}</th>"

    html += "</tr>"

    for idx, row in df.iterrows():
        html += f"<tr><td><b>{idx}</b></td>"

        for val in row:
            val = clean_team(val)
            color = color_map.get(val, "#ffffff")
            text_color = "white" if val in ["LEAVE", "NON-WORK DAY"] else "black"

            html += (
            f"<td style='background-color:{color};color:{text_color}'>"
            f"{val}</td>"
            )

        html += "</tr>"

    html += "</table>"

    st.markdown(html, unsafe_allow_html=True)


# =========================
# LEGEND
# =========================
def render_legend(color_map):
    st.sidebar.subheader("Legend")

    for key, color in color_map.items():
        st.sidebar.markdown(
            f"<div style='display:flex;align-items:center;margin-bottom:5px;'>"
            f"<div style='width:14px;height:14px;background:{color};margin-right:8px;'></div>"
            f"{key}</div>",
            unsafe_allow_html=True
        )


# =========================
# MAIN APP
# =========================
movement, lifecycle, resources, leaves = load_data()
names, holidays = load_lists()

st.title("Workforce Timeline System")

mode = st.selectbox("Mode", ["A", "D"])

start_date = st.date_input("START DATE")
end_date = st.date_input("END DATE")

timeline = build_timeline(start_date, end_date)

# =========================
# MODE A: EMPLOYEE VIEW
# =========================
if mode == "A":
    name_input = st.text_input("Employee Search")

    matched = [
        n for n in names
        if isinstance(n, str) and name_input.lower() in n.lower()
    ]

    data = {}

    for emp in matched:
        data[emp] = resolve_employee(emp, movement, timeline, holidays, leaves)

    if data:
        df = pd.DataFrame(data, index=timeline).T
        df = df.sort_index()

        all_values = []
        for row in data.values():
            all_values.extend([v for v in row])

        color_map = generate_color_map(all_values)

        st.subheader("Employee Timeline")
        render_colored_table(df, color_map)
        render_legend(color_map)

# =========================
# MODE D: TEAM → EMPLOYEES
# =========================
elif mode == "D":
    team_options = get_all_core_teams(movement)
    team_input = st.selectbox("Team / Client Name", team_options)

    if team_input:

        team_input = team_input.strip()

        # =========================
        # FIND EMPLOYEES BY CORE TEAM
        # =========================
        matched = []

        for emp in names:
            if not isinstance(emp, str):
                continue

            core_team = get_core_team(emp, movement)

            if core_team == team_input:
                matched.append(emp)

        st.subheader(f"Team Timeline: {team_input}")

        if not matched:
            st.warning("No employees found for this team/client.")
        else:

            # =========================
            # SPLIT BY ROLE (PL / PA)
            # =========================
            pl_list = []
            pa_list = []

            for emp in matched:
                role = resources.loc[resources["NAME"] == emp, "CURRENT PROD ROLE"]
                role = role.iloc[0] if not role.empty else None

                if role == "PL":
                    pl_list.append(emp)
                elif role == "PA":
                    pa_list.append(emp)

            # =========================
            # FUNCTION TO BUILD TABLE
            # =========================
            def build_table(group):
                data = {}

                for emp in group:
                    data[emp] = resolve_employee(emp, movement, timeline, holidays, leaves)

                if not data:
                    return None, None

                df = pd.DataFrame(data, index=timeline).T
                df = df.sort_index()

                all_values = []
                for row in data.values():
                    all_values.extend([v for v in row])

                color_map = generate_color_map(all_values)

                return df, color_map

            # =========================
            # PL TABLE
            # =========================
            if pl_list:
                st.subheader("PL")

                df_pl, color_map_pl = build_table(pl_list)

                if df_pl is not None:
                    render_colored_table(df_pl, color_map_pl)
                    render_legend(color_map_pl)

            # =========================
            # PA TABLE
            # =========================
            if pa_list:
                st.subheader("PA")

                df_pa, color_map_pa = build_table(pa_list)

                if df_pa is not None:
                    render_colored_table(df_pa, color_map_pa)
                    render_legend(color_map_pa)

            # =========================
            # SUMMARY LIST (OPTIONAL)
            # =========================
            st.subheader("Employees in Team")

            st.dataframe(pd.DataFrame(matched, columns=["EMPLOYEES"]))
